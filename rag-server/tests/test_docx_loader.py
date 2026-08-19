"""Tests for DOCX document loader."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from docx import Document

from app.ingestion.docx_loader import DocxLoadError, DocxLoader


class TestDocxZipValidation:
    """Test DOCX ZIP container validation."""

    def test_invalid_zip_rejected(self, tmp_path: Path) -> None:
        """Test rejection of invalid ZIP files."""
        loader = DocxLoader()

        docx_file = tmp_path / "invalid.docx"
        docx_file.write_bytes(b"NOT_A_ZIP")

        with pytest.raises(DocxLoadError) as exc_info:
            loader.load(docx_file, "invalid.docx", "hash")
        assert "invalid ZIP" in str(exc_info.value)

    def test_missing_content_types_rejected(self, tmp_path: Path) -> None:
        """Test rejection of DOCX missing [Content_Types].xml."""
        loader = DocxLoader()

        docx_file = tmp_path / "no_types.docx"

        # Create minimal invalid DOCX
        with zipfile.ZipFile(docx_file, "w") as zf:
            zf.writestr("word/document.xml", "<document/>")

        with pytest.raises(DocxLoadError) as exc_info:
            loader.load(docx_file, "no_types.docx", "hash")
        assert "missing [Content_Types].xml" in str(exc_info.value)

    def test_missing_document_xml_rejected(self, tmp_path: Path) -> None:
        """Test rejection of DOCX missing word/document.xml."""
        loader = DocxLoader()

        docx_file = tmp_path / "no_doc.docx"

        with zipfile.ZipFile(docx_file, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")

        with pytest.raises(DocxLoadError) as exc_info:
            loader.load(docx_file, "no_doc.docx", "hash")
        assert "missing word/document.xml" in str(exc_info.value)

    def test_absolute_archive_path_rejected(self, tmp_path: Path) -> None:
        """Test rejection of absolute paths in archive."""
        loader = DocxLoader()

        docx_file = tmp_path / "absolute_path.docx"

        with zipfile.ZipFile(docx_file, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")
            zf.writestr("word/document.xml", "<document/>")
            zf.writestr("/etc/passwd", "hacked")

        with pytest.raises(DocxLoadError) as exc_info:
            loader.load(docx_file, "absolute_path.docx", "hash")
        assert "absolute archive path" in str(exc_info.value)

    def test_traversal_in_archive_rejected(self, tmp_path: Path) -> None:
        """Test rejection of .. traversal in archive."""
        loader = DocxLoader()

        docx_file = tmp_path / "traversal.docx"

        with zipfile.ZipFile(docx_file, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")
            zf.writestr("word/document.xml", "<document/>")
            zf.writestr("../../../etc/passwd", "hacked")

        with pytest.raises(DocxLoadError) as exc_info:
            loader.load(docx_file, "traversal.docx", "hash")
        assert "archive traversal" in str(exc_info.value)

    def test_excessive_entries_rejected(self, tmp_path: Path) -> None:
        """Test rejection of excessive ZIP entries."""
        loader = DocxLoader(max_entries=10)

        docx_file = tmp_path / "too_many.docx"

        with zipfile.ZipFile(docx_file, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")
            zf.writestr("word/document.xml", "<document/>")

            # Add many entries
            for i in range(15):
                zf.writestr(f"file{i}.xml", "<data/>")

        with pytest.raises(DocxLoadError) as exc_info:
            loader.load(docx_file, "too_many.docx", "hash")
        assert "too many ZIP entries" in str(exc_info.value)

    def test_excessive_uncompressed_size_rejected(self, tmp_path: Path) -> None:
        """Test rejection of excessive uncompressed size."""
        loader = DocxLoader(max_uncompressed_bytes=1024)

        docx_file = tmp_path / "too_large.docx"

        with zipfile.ZipFile(docx_file, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")
            zf.writestr("word/document.xml", "<document/>")
            # Add large data
            zf.writestr("data.xml", "x" * 2048)

        with pytest.raises(DocxLoadError) as exc_info:
            loader.load(docx_file, "too_large.docx", "hash")
        assert "uncompressed size exceeds limit" in str(exc_info.value)

    def test_docm_macro_enabled_rejected(self, tmp_path: Path) -> None:
        """Test rejection of macro-enabled DOCM files."""
        loader = DocxLoader()

        docx_file = tmp_path / "macro.docm"

        with zipfile.ZipFile(docx_file, "w") as zf:
            # Content types with macro reference
            zf.writestr(
                "[Content_Types].xml",
                '<Types><Override PartName="/word/macros.xml"/></Types>',
            )
            zf.writestr("word/document.xml", "<document/>")

        with pytest.raises(DocxLoadError) as exc_info:
            loader.load(docx_file, "macro.docm", "hash")
        assert "macro" in str(exc_info.value).lower()


class TestDocxContentExtraction:
    """Test DOCX content extraction."""

    def test_heading_extracted_as_markdown(self, tmp_path: Path) -> None:
        """Test that Word headings are converted to Markdown."""
        docx_file = tmp_path / "headings.docx"

        # Create DOCX dynamically
        doc = Document()
        doc.add_heading("Main Title", level=1)
        doc.add_paragraph("Some text")
        doc.add_heading("Subheading", level=2)
        doc.add_paragraph("More text")
        doc.save(docx_file)

        loader = DocxLoader()
        loaded = loader.load(docx_file, "headings.docx", "hash")

        # Check markdown headings
        assert "# Main Title" in loaded.content
        assert "## Subheading" in loaded.content

    def test_paragraph_count_recorded(self, tmp_path: Path) -> None:
        """Test that paragraph count is recorded."""
        docx_file = tmp_path / "paragraphs.docx"

        doc = Document()
        doc.add_paragraph("Para 1")
        doc.add_paragraph("Para 2")
        doc.add_paragraph("Para 3")
        doc.save(docx_file)

        loader = DocxLoader()
        loaded = loader.load(docx_file, "paragraphs.docx", "hash")

        assert loaded.metadata["paragraph_count"] == 3

    def test_table_extracted_as_markdown(self, tmp_path: Path) -> None:
        """Test that tables are converted to Markdown."""
        docx_file = tmp_path / "table.docx"

        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        table.rows[0].cells[0].text = "Header 1"
        table.rows[0].cells[1].text = "Header 2"
        table.rows[0].cells[2].text = "Header 3"
        table.rows[1].cells[0].text = "Data 1"
        table.rows[1].cells[1].text = "Data 2"
        table.rows[1].cells[2].text = "Data 3"
        doc.save(docx_file)

        loader = DocxLoader()
        loaded = loader.load(docx_file, "table.docx", "hash")

        # Check Markdown table format
        assert "|" in loaded.content
        assert "---" in loaded.content
        assert "Header 1" in loaded.content
        assert "Data 1" in loaded.content

    def test_table_count_recorded(self, tmp_path: Path) -> None:
        """Test that table count is recorded."""
        docx_file = tmp_path / "tables.docx"

        doc = Document()
        doc.add_table(rows=2, cols=2)
        doc.add_table(rows=3, cols=3)
        doc.save(docx_file)

        loader = DocxLoader()
        loaded = loader.load(docx_file, "tables.docx", "hash")

        assert loaded.metadata["table_count"] == 2

    def test_pipe_characters_escaped_in_tables(self, tmp_path: Path) -> None:
        """Test that pipe characters are escaped in table cells."""
        docx_file = tmp_path / "pipe_table.docx"

        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Header"
        table.rows[0].cells[1].text = "Header"
        table.rows[1].cells[0].text = "Data | with pipe"
        table.rows[1].cells[1].text = "Data"
        doc.save(docx_file)

        loader = DocxLoader()
        loaded = loader.load(docx_file, "pipe_table.docx", "hash")

        # Check that pipes are escaped
        assert "\\|" in loaded.content


class TestDocxEmptyRejection:
    """Test rejection of empty documents."""

    def test_empty_docx_rejected(self, tmp_path: Path) -> None:
        """Test rejection of DOCX with no extractable text."""
        docx_file = tmp_path / "empty.docx"

        doc = Document()
        # Don't add any content
        doc.save(docx_file)

        loader = DocxLoader()

        with pytest.raises(DocxLoadError) as exc_info:
            loader.load(docx_file, "empty.docx", "hash")
        assert "empty_document" in str(exc_info.value)


class TestDocxMetadata:
    """Test DOCX metadata extraction."""

    def test_file_sha256_recorded(self, tmp_path: Path) -> None:
        """Test that file SHA-256 is recorded."""
        docx_file = tmp_path / "test.docx"

        doc = Document()
        doc.add_paragraph("Test content")
        doc.save(docx_file)

        loader = DocxLoader()
        test_hash = "abc123def456"
        loaded = loader.load(docx_file, "test.docx", test_hash)

        assert loaded.metadata["file_sha256"] == test_hash

    def test_title_from_filename(self, tmp_path: Path) -> None:
        """Test that title comes from filename when no core property."""
        docx_file = tmp_path / "my_document.docx"

        doc = Document()
        doc.add_paragraph("Content")
        doc.save(docx_file)

        loader = DocxLoader()
        loaded = loader.load(docx_file, "my_document.docx", "hash")

        assert loaded.title == "my_document"

    def test_title_from_core_properties(self, tmp_path: Path) -> None:
        """Test that title from core properties takes precedence."""
        docx_file = tmp_path / "filename.docx"

        doc = Document()
        doc.core_properties.title = "Custom Title"
        doc.add_paragraph("Content")
        doc.save(docx_file)

        loader = DocxLoader()
        loaded = loader.load(docx_file, "filename.docx", "hash")

        assert loaded.title == "Custom Title"


class TestDocxSourceKey:
    """Test source key generation."""

    def test_source_key_format(self, tmp_path: Path) -> None:
        """Test that source key has document: prefix."""
        docx_file = tmp_path / "doc.docx"

        doc = Document()
        doc.add_paragraph("Content")
        doc.save(docx_file)

        loader = DocxLoader()
        loaded = loader.load(docx_file, "doc.docx", "hash")

        assert loaded.source_key == "document:doc.docx"

    def test_source_type_is_project_document(self, tmp_path: Path) -> None:
        """Test that source_type is project_document."""
        docx_file = tmp_path / "doc.docx"

        doc = Document()
        doc.add_paragraph("Content")
        doc.save(docx_file)

        loader = DocxLoader()
        loaded = loader.load(docx_file, "doc.docx", "hash")

        assert loaded.source_type == "project_document"


class TestDocxLineNormalization:
    """Test line ending normalization."""

    def test_blank_lines_collapsed(self, tmp_path: Path) -> None:
        """Test that excessive blank lines are collapsed."""
        docx_file = tmp_path / "blanks.docx"

        doc = Document()
        doc.add_paragraph("Para 1")
        doc.add_paragraph("")
        doc.add_paragraph("")
        doc.add_paragraph("")
        doc.add_paragraph("Para 2")
        doc.save(docx_file)

        loader = DocxLoader()
        loaded = loader.load(docx_file, "blanks.docx", "hash")

        # Check that excessive blank lines are collapsed
        assert "\n\n\n" not in loaded.content
